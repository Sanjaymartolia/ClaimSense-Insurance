"""Builds the team guide in three formats from one source, docs/TEAM_GUIDE.md.

python scripts/build_team_guide.py   (needs: pip install -e ".[docs]", and Edge or Chrome)
  -> docs/diagrams/team-workflow.png + .svg   the workflow figure (edit LANES below)
  -> docs/ClaimSense_Team_Guide.html + .pdf   via build_pdf.py's converter and the browser
  -> docs/ClaimSense_Team_Guide.docx          via python-docx
Edit the Markdown, re-run, and all three stay identical.
"""

import re
import subprocess

import build_pdf as pdf
import matplotlib
from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from matplotlib.patches import FancyBboxPatch  # noqa: E402

DOCS = pdf.DOCS
SRC = DOCS / "TEAM_GUIDE.md"
OUT = DOCS / "ClaimSense_Team_Guide"
FIG = DOCS / "diagrams" / "team-workflow"
NAVY = RGBColor(0x1F, 0x38, 0x64)

# ------------------------------------------------------------------ the figure --
# (lane title, colour, fill, [step, ...], loops back to its first step?)
LANES = [
    (
        "1  Everyone, once",
        "#1971c2",
        "#e7f5ff",
        [
            "Clone repo",
            'venv +\npip install -e\n".[dev,agents,\nretrieval,course]"',
            "Key in .env\n(repo root,\nnever committed)",
            "new_student.py\n<you>\n-> exercises/<you>/",
            "build_index.py\n--force\n+ pytest tests",
        ],
        False,
    ),
    (
        "2  Course-kit loop\n    after every TODO",
        "#2b8a3e",
        "#ebfbee",
        [
            "Pick next TODO\n4 > 2 > 3 > 6 > 5 > 1",
            "Write code\n(hints are in\nthe comments)",
            "Contract tests\npytest tests\n(no LLM, free)",
            "evaluate_dev.py\n--label X [--redteam]\nDev 20 · Red-team 30\nRetrieval 12",
            "Read the plot\nruns/log.png\n+ worst claims",
            "Bug? failing test\nin test_my_bugs.py\nthen fix",
            "REPORT.md\nnumbers + plots",
        ],
        True,
    ),
    (
        "3  Product work\n    1 task = 1 PR",
        "#5f3dc4",
        "#f3f0ff",
        [
            "Task Txx\n(tracker)",
            "Branch +\nplan mode",
            "Tests first\n(hand-worked\nfor calc/)",
            "Implement in\nyour role's\nfolder",
            "ruff + pytest\nlocally",
            "Small PR\nreview by\nanother role",
            "CI green\n-> merge main",
        ],
        False,
    ),
    (
        "4  Eval & Red-Team\n    once, frozen",
        "#c92a2a",
        "#ffe3e3",
        [
            "Freeze code,\nmodel, prompts,\nindex",
            "run_batch.py\nheld-out 60\n(outside repo)",
            "score_traces.py\n+ plot",
            "Report result\nnever tune on it",
        ],
        False,
    ),
]


def draw_figure():
    fig, ax = plt.subplots(figsize=(17, 9.6))
    ax.set_xlim(0, 100)
    ax.set_ylim(0, 58)
    ax.axis("off")
    ax.text(
        50,
        56.5,
        "ClaimSense team workflow: who runs what, and when",
        ha="center",
        fontsize=15,
        weight="bold",
        color="#1f3864",
    )
    left, right, lane_h = 15.5, 99.0, 12.5
    for row, (title, edge, fill, steps, loops) in enumerate(LANES):
        y0 = 54 - (row + 1) * (lane_h + 0.6)
        ax.add_patch(
            FancyBboxPatch(
                (0.5, y0),
                99,
                lane_h,
                boxstyle="round,pad=0,rounding_size=0.8",
                fc="#fbfcfe",
                ec="#ced4da",
                lw=1,
            )
        )
        ax.text(1.5, y0 + lane_h / 2, title, va="center", fontsize=10.5, weight="bold", color=edge)
        gap = 1.1
        w = (right - left - gap * (len(steps) - 1)) / len(steps)
        h, yb = 8.4, y0 + (lane_h - 8.4) / 2
        for i, step in enumerate(steps):
            x = left + i * (w + gap)
            ax.add_patch(
                FancyBboxPatch(
                    (x, yb),
                    w,
                    h,
                    boxstyle="round,pad=0,rounding_size=0.6",
                    fc=fill,
                    ec=edge,
                    lw=1.4,
                )
            )
            ax.text(
                x + w / 2,
                yb + h / 2,
                step,
                ha="center",
                va="center",
                fontsize=8.1,
                color="#212529",
                linespacing=1.15,
            )
            if i:
                ax.annotate(
                    "",
                    (x, yb + h / 2),
                    (x - gap, yb + h / 2),
                    arrowprops={"arrowstyle": "->", "color": edge, "lw": 1.4},
                )
        if loops:  # the loop arrow: from the plot / bug step back to "Pick next TODO"
            # drawn as an explicit path that stays inside the lane, under the boxes
            x_from, x_to, y_low = left + 5 * (w + gap) + w / 2, left + w / 2, yb - 1.0
            ax.plot([x_from, x_from, x_to], [yb, y_low, y_low], color=edge, lw=1.4)
            ax.annotate(
                "",
                (x_to, yb),
                (x_to, y_low),
                arrowprops={"arrowstyle": "->", "color": edge, "lw": 1.4},
            )
            ax.text(
                (x_to + x_from) / 2,
                y_low,
                "repeat for the next TODO",
                ha="center",
                va="center",
                fontsize=8,
                style="italic",
                color=edge,
                bbox={"fc": "#fbfcfe", "ec": "none", "pad": 1},
            )
    ax.text(
        50,
        0.4,
        "Rules in every lane: model proposes, code decides, a person approves · "
        "least code that works · bug = failing test first · plot before any claim",
        ha="center",
        fontsize=9,
        color="#495057",
    )
    FIG.parent.mkdir(parents=True, exist_ok=True)
    for ext in ("png", "svg"):
        fig.savefig(FIG.with_suffix(f".{ext}"), dpi=200, bbox_inches="tight")
    plt.close(fig)


# ---------------------------------------------------------------------- the PDF --
def build_pdf():
    html_out = OUT.with_suffix(".html")
    html_out.write_text(
        f'<!doctype html><html><head><meta charset="utf-8"><title>ClaimSense Team Guide</title>'
        f"<style>{pdf.CSS}</style></head><body>{pdf.convert(SRC.read_text('utf-8'))}</body></html>",
        encoding="utf-8",
    )
    subprocess.run(
        [
            pdf.find_browser(),
            "--headless",
            "--disable-gpu",
            "--no-pdf-header-footer",
            f"--print-to-pdf={OUT.with_suffix('.pdf')}",
            html_out.as_uri(),
        ],
        check=True,
        capture_output=True,
        timeout=120,
    )


# --------------------------------------------------------------------- the DOCX --
def shade(element, fill):
    """Background colour for a table cell or a paragraph (code blocks)."""
    pr = (
        element.get_or_add_tcPr()
        if hasattr(element, "get_or_add_tcPr")
        else element.get_or_add_pPr()
    )
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pr.append(shd)


def add_inline(par, text, size=None, color=None):
    """**bold**, *italic* and `code` -> runs."""
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`|(?<![*\w])\*[^*]+\*(?!\*))", text):
        if not part:
            continue
        if part.startswith("**"):
            run = par.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`"):
            run = par.add_run(part[1:-1])
            run.font.name, run.font.color.rgb = "Consolas", RGBColor(0xC9, 0x2A, 0x2A)
        elif part.startswith("*"):
            run = par.add_run(part[1:-1])
            run.italic = True
        else:
            run = par.add_run(part)
        if size:
            run.font.size = size
        if color:
            run.font.color.rgb = color


def build_docx():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(1.9))
    doc.styles["Normal"].font.name, doc.styles["Normal"].font.size = "Calibri", Pt(10.5)
    lines, i = SRC.read_text("utf-8").splitlines(), 0
    while i < len(lines):
        line = lines[i]
        if not line.strip():
            i += 1
        elif line.startswith("    "):  # code block: one shaded paragraph, a line break per line
            block = []
            while i < len(lines) and (lines[i].startswith("    ") or not lines[i].strip()):
                block.append(lines[i][4:])
                i += 1
            par = doc.add_paragraph()
            shade(par._p, "F1F3F5")
            for n, text in enumerate("\n".join(block).rstrip().split("\n")):
                run = par.add_run(text)
                run.font.name, run.font.size = "Consolas", Pt(8.5)
                if n < len("\n".join(block).rstrip().split("\n")) - 1:
                    run.add_break(WD_BREAK.LINE)
        elif m := re.match(r"(#{1,3}) (.*)", line):
            level = len(m.group(1)) - 1  # '#' -> Title, '##' -> Heading 1, '###' -> Heading 2
            heading = doc.add_heading(m.group(2).replace("`", ""), level)
            for run in heading.runs:
                run.font.color.rgb = NAVY
            i += 1
        elif line.strip() == "---":
            i += 1
        elif m := re.match(r"!\[[^\]]*\]\(([^)]+)\)", line):
            doc.add_picture(str(DOCS / m.group(1)), width=Cm(17.2))
            i += 1
        elif line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            rows = [rows[0]] + rows[2:]
            table = doc.add_table(rows=len(rows), cols=len(rows[0]), style="Table Grid")
            for r, cells in enumerate(rows):
                for c, text in enumerate(cells):
                    cell = table.cell(r, c)
                    cell.paragraphs[0].text = ""
                    if r == 0:
                        shade(cell._tc, "1F3864")
                    add_inline(
                        cell.paragraphs[0],
                        text,
                        Pt(8.5),
                        RGBColor(0xFF, 0xFF, 0xFF) if r == 0 else None,
                    )
                    if r == 0:
                        for run in cell.paragraphs[0].runs:
                            run.bold = True
            doc.add_paragraph()
        elif re.match(r"(\d+\.|-) ", line):
            n = 0
            while i < len(lines) and (
                re.match(r"(\d+\.|-) ", lines[i])
                or (lines[i].startswith("   ") and lines[i].strip())
            ):
                if re.match(r"(\d+\.|-) ", lines[i]):
                    n += 1
                    ordered = lines[i][0].isdigit()
                    par = doc.add_paragraph(style="List Bullet" if not ordered else None)
                    if ordered:  # manual numbers: Word's list numbering runs on across lists
                        par.paragraph_format.left_indent = Cm(0.6)
                        par.add_run(f"{n}. ")
                    add_inline(par, re.sub(r"^(\d+\.|-) ", "", lines[i]))
                else:
                    add_inline(par, " " + lines[i].strip())
                i += 1
        else:
            text = []
            while (
                i < len(lines)
                and lines[i].strip()
                and not re.match(r"(#{1,3} |\||---$|    |(\d+\.|-) |!\[)", lines[i])
            ):
                text.append(lines[i].strip())
                i += 1
            add_inline(doc.add_paragraph(), " ".join(text))
    doc.save(OUT.with_suffix(".docx"))


if __name__ == "__main__":
    draw_figure()
    build_pdf()
    build_docx()
    print("wrote", FIG.with_suffix(".png"), OUT.with_suffix(".pdf"), OUT.with_suffix(".docx"))
